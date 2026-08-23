"""
extract_transcript.py

Google Meets が生成した .docx 議事録(Geminiによるメモ)を読み込み、
話者付きの発言リスト(JSON)に変換する。

議事録の構造:
  - メタデータ (タイトル、日付、参加者) は冒頭部分
  - 「概要」「次のステップ」「詳細」はサマリセクション(分析対象外)
  - 「📖 文字起こし」以降が分析対象の発言データ
    - Heading 3 スタイル = タイムスタンプ (00:00:00)
    - normal スタイル = 発言ブロック
      - bold run = 話者名 (末尾コロン付き)
      - 通常 run = 発言内容

実行例:
    python scripts/extract_transcript.py \
        --input /path/to/meeting.docx \
        --output /path/to/output.json
"""

import argparse
import json
import re
import sys
from dataclasses import dataclass, asdict
from datetime import datetime
from pathlib import Path

from docx import Document


@dataclass
class Utterance:
    """単一の発言を表すデータクラス"""
    speaker: str
    timestamp: str
    text: str
    meeting_title: str
    meeting_date: str
    section_index: int
    raw_index: int


def extract_meeting_metadata(doc) -> dict:
    """議事録の冒頭からメタデータを抽出"""
    meta = {
        "title": "",
        "date": "",
        "participants": [],
    }
    
    for para in doc.paragraphs[:50]:
        text = para.text.strip()
        if not text:
            continue
        
        style_name = para.style.name if para.style else ""
        
        if not meta["title"]:
            if style_name == "Heading 2" or text.startswith("##"):
                if "文字起こし" not in text:
                    meta["title"] = text.lstrip("#").strip()
        
        date_match = re.search(r"(\d{1,2})月\s*(\d{1,2}),?\s*(\d{4})", text)
        if date_match and not meta["date"]:
            month, day, year = date_match.groups()
            meta["date"] = f"{year}-{int(month):02d}-{int(day):02d}"
    
    return meta


def extract_utterances(doc, meta: dict) -> list[Utterance]:
    """議事録から発言を抽出"""
    utterances = []
    current_timestamp = "00:00:00"
    section_index = 0
    raw_index = 0
    in_transcript = False
    
    for para in doc.paragraphs:
        text = para.text.strip()
        style_name = para.style.name if para.style else ""
        
        # 「📖 文字起こし」セクションに入ったら開始
        if not in_transcript:
            if "文字起こし" in text and (style_name == "Title" or
                                         "📖" in text or
                                         style_name == "Heading 2"):
                in_transcript = True
            continue
        
        # 文字起こし終了マーカー
        if "文字起こしが終了" in text or "コンピュータが生成" in text:
            break
        
        # タイムスタンプ検出 (Heading 3 スタイル、HH:MM:SS のパターン)
        if style_name == "Heading 3":
            ts_match = re.search(r"(\d{2}:\d{2}:\d{2})", text)
            if ts_match:
                current_timestamp = ts_match.group(1)
                section_index += 1
                continue
        
        # 発言ブロック処理 (normal スタイル、bold run が話者名)
        if style_name == "normal" and para.runs:
            current_speaker = None
            current_text_buffer = []
            
            for run in para.runs:
                run_text = run.text
                if not run_text:
                    continue
                
                if run.bold:
                    # 前の話者の発言をフラッシュ
                    if current_speaker and current_text_buffer:
                        utterance_text = " ".join(current_text_buffer).strip()
                        utterance_text = re.sub(r"\s+", " ", utterance_text)
                        if utterance_text:
                            utterances.append(Utterance(
                                speaker=current_speaker,
                                timestamp=current_timestamp,
                                text=utterance_text,
                                meeting_title=meta.get("title", ""),
                                meeting_date=meta.get("date", ""),
                                section_index=section_index,
                                raw_index=raw_index,
                            ))
                            raw_index += 1
                    
                    current_speaker = run_text.rstrip(":：").strip()
                    current_text_buffer = []
                else:
                    if current_speaker:
                        current_text_buffer.append(run_text)
            
            # 最後の話者の発言をフラッシュ
            if current_speaker and current_text_buffer:
                utterance_text = " ".join(current_text_buffer).strip()
                utterance_text = re.sub(r"\s+", " ", utterance_text)
                if utterance_text:
                    utterances.append(Utterance(
                        speaker=current_speaker,
                        timestamp=current_timestamp,
                        text=utterance_text,
                        meeting_title=meta.get("title", ""),
                        meeting_date=meta.get("date", ""),
                        section_index=section_index,
                        raw_index=raw_index,
                    ))
                    raw_index += 1
    
    return utterances


def merge_consecutive_utterances(utterances: list[Utterance]) -> list[Utterance]:
    """同じ話者の連続する発言を統合"""
    if not utterances:
        return []
    
    merged = [utterances[0]]
    
    for u in utterances[1:]:
        last = merged[-1]
        if u.speaker == last.speaker and u.section_index == last.section_index:
            merged[-1] = Utterance(
                speaker=last.speaker,
                timestamp=last.timestamp,
                text=last.text + " " + u.text,
                meeting_title=last.meeting_title,
                meeting_date=last.meeting_date,
                section_index=last.section_index,
                raw_index=last.raw_index,
            )
        else:
            merged.append(u)
    
    return merged


def extract_transcript(input_path: str) -> dict:
    """メイン処理"""
    doc = Document(input_path)
    
    meta = extract_meeting_metadata(doc)
    utterances = extract_utterances(doc, meta)
    utterances = merge_consecutive_utterances(utterances)
    
    return {
        "meta": {
            "title": meta["title"],
            "date": meta["date"],
            "participants": meta["participants"],
            "source_file": str(Path(input_path).name),
            "total_utterances": len(utterances),
            "extracted_at": datetime.now().isoformat(),
        },
        "utterances": [asdict(u) for u in utterances],
    }


def main():
    parser = argparse.ArgumentParser(
        description="Google Meets議事録から発言を抽出"
    )
    parser.add_argument("--input", "-i", required=True, help="入力 .docx ファイルのパス")
    parser.add_argument("--output", "-o", required=True, help="出力 JSON ファイルのパス")
    parser.add_argument("--pretty", action="store_true", help="JSON を整形して出力")
    
    args = parser.parse_args()
    
    if not Path(args.input).exists():
        print(f"エラー: 入力ファイルが見つかりません: {args.input}", file=sys.stderr)
        sys.exit(1)
    
    try:
        result = extract_transcript(args.input)
    except Exception as e:
        print(f"エラー: 議事録の抽出に失敗しました: {e}", file=sys.stderr)
        sys.exit(1)
    
    Path(args.output).parent.mkdir(parents=True, exist_ok=True)
    
    with open(args.output, "w", encoding="utf-8") as f:
        if args.pretty:
            json.dump(result, f, ensure_ascii=False, indent=2)
        else:
            json.dump(result, f, ensure_ascii=False)
    
    print(f"抽出完了: {result['meta']['total_utterances']} 発言を {args.output} に保存",
          file=sys.stderr)


if __name__ == "__main__":
    main()
